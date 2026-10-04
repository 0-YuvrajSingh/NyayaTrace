import copy, hashlib, os, re, shutil
from docx import Document

SRC = "/repo/Indian_Legal_XAI.docx"
OUT_DIR = "/repo/experiments/audit_fixes/spec_clean"
OUT = os.path.join(OUT_DIR, "Indian_Legal_XAI_clean.docx")

def sha(p):
    with open(p, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()

def set_text(par, text):
    """Replace the whole paragraph text, keeping the first run's formatting."""
    if par.runs:
        par.runs[0].text = text
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.add_run(text)

def replace_in_par(par, old, new):
    """Run-level replace first (keeps bold etc.); fall back to joined text."""
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    if old in par.text:
        set_text(par, par.text.replace(old, new))
        return True
    return False

def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    before = sha(SRC)
    shutil.copyfile(SRC, OUT)
    doc = Document(OUT)
    log = []

    # 1-2. Title and subtitle
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Canonical Project Specification"):
            set_text(p, "Canonical Project Specification (Final)")
            log.append("title replaced")
        elif t.startswith("Supersedes prior working drafts"):
            set_text(p, "14\u201316 week semester scope")
            log.append("subtitle replaced")

    # 3-4. Status box: replace row 1, delete the companion-synopsis row
    for tbl in doc.tables:
        first = tbl.rows[0].cells[0].text.strip()
        if first.startswith("Document status and governing rule"):
            cell = tbl.rows[1].cells[0]
            set_text(cell.paragraphs[0],
                     "This is the canonical and final project specification. "
                     "Implementation and research follow it. Changes are governed "
                     "by the Scope Freeze and Change Control section.")
            for extra in cell.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
            log.append("status box text replaced")
            for row in list(tbl.rows[2:]):
                if row.cells[0].text.strip().startswith("A companion full-length synopsis"):
                    tbl._tbl.remove(row._tr)
                    log.append("companion synopsis row deleted")

    # 5. Evaluation Metrics scope note
    for p in doc.paragraphs:
        if replace_in_par(p, "not part of the approved baseline",
                          "not part of the approved evaluation set"):
            log.append("scope note reworded")

    # 6. Amendment A-1 directly under the decision_date <= case_date line
    amended = False
    for p in doc.paragraphs:
        if "decision_date <= case_date" in p.text and not amended:
            new_p = copy.deepcopy(p._p)
            p._p.addnext(new_p)
            from docx.text.paragraph import Paragraph
            np_ = Paragraph(new_p, p._parent)
            set_text(np_,
                "Amendment A-1 (implementation, dated 28 September 2026): ILDC "
                "supplies only the year of each query case, not its exact decision "
                "date. Precedent eligibility is therefore implemented as "
                "decision_year < query_year. Same-year authorities and authorities "
                "with missing or unparseable dates are excluded. This is stricter "
                "than decision_date <= case_date and cannot admit a later "
                "authority; recall for same-year authorities may be lower. See "
                "experiments/audit_fixes/TEMPORAL_ERRATUM.md.")
            amended = True
            log.append("Amendment A-1 inserted")

    doc.save(OUT)

    # Verification: leftovers and original untouched
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                texts.append(c.text)
    leftovers = [t[:90] for t in texts if re.search(r"Document [123]\b", t)]
    print("CHANGES:", *log, sep="\n  - ")
    print("LEFTOVER 'Document N' references:", leftovers or "none")
    print("<= line still present:", any("decision_date <= case_date" in t for t in texts))
    print("original unchanged:", before == sha(SRC))
    print("clean copy SHA-256:", sha(OUT))

if __name__ == "__main__":
    main()
